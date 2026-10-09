"""Owned test artwork. Requires python-pptx, Pillow and LibreOffice."""
from pathlib import Path
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw
import subprocess
import tempfile

output = Path(__file__).resolve().parent
with tempfile.TemporaryDirectory() as temporary:
    temporary = Path(temporary)
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    def title(slide, value):
        box = slide.shapes.add_textbox(Inches(.6), Inches(.5), Inches(9), Inches(1))
        box.text = value
        box.text_frame.paragraphs[0].font.size = Pt(30)
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    title(first, 'LOCAL FIELD NOTES')
    text = first.shapes.add_textbox(Inches(.6), Inches(1.6), Inches(8.8), Inches(1))
    text.text = 'A legacy presentation, read in your browser.\nText, shapes, and pictures — three saved slides.'
    text.text_frame.paragraphs[0].font.size = Pt(22)
    shape = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(3), Inches(3), Inches(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(24, 110, 91)
    shape.text = '01 / OBSERVE'
    image = Image.new('RGB', (600, 400), '#e9dbc3')
    draw = ImageDraw.Draw(image)
    draw.ellipse((160, 30, 480, 350), fill='#b96e48')
    draw.rectangle((0, 270, 600, 400), fill='#36594d')
    data = BytesIO()
    image.save(data, 'PNG')
    first.shapes.add_picture(BytesIO(data.getvalue()), Inches(4.2), Inches(3), Inches(4), Inches(2.66))
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    title(second, 'CHECK THE DETAILS')
    table = second.shapes.add_table(3, 2, Inches(.6), Inches(1.8), Inches(8.5), Inches(2)).table
    for row, values in enumerate([['Material', 'Observation'], ['Paper', 'Warm ivory'], ['Ink', 'Forest green']]):
        for column, value in enumerate(values): table.cell(row, column).text = value
    rotated = second.shapes.add_textbox(Inches(1), Inches(4.4), Inches(3), Inches(1))
    rotated.text = 'Rotated annotation'
    rotated.rotation = 15
    third = presentation.slides.add_slide(presentation.slide_layouts[6])
    title(third, 'ONE MORE LOOK')
    third.shapes.add_textbox(Inches(.6), Inches(2), Inches(8), Inches(1)).text = 'Use the overview to return to slide one.'
    source = temporary / 'field-notes.pptx'
    presentation.save(source)
    subprocess.run(['libreoffice', '-env:UserInstallation=file://' + str(temporary / 'profile'), '--headless', '--convert-to', 'ppt:MS PowerPoint 97', '--outdir', str(output), str(source)], check=True)
