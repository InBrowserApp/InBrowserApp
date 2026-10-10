"""Generate owned interoperability fixtures with openpyxl; no external documents."""
from pathlib import Path
from datetime import datetime, timedelta, time
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

root = Path(__file__).parent
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def caches(path, values):
    with ZipFile(path) as z:
        files = {name: z.read(name) for name in z.namelist()}
    xml = ET.fromstring(files['xl/worksheets/sheet1.xml'])
    for cell in xml.findall('.//s:c', ns):
        if cell.attrib['r'] not in values:
            continue
        kind, value = values[cell.attrib['r']]
        cell.set('t', kind)
        cell.find('s:v', ns).text = value
    # Keep default namespace: this is a valid ordinary OOXML worksheet.
    ET.register_namespace('', ns['s'])
    files['xl/worksheets/sheet1.xml'] = ET.tostring(xml, encoding='utf-8', xml_declaration=True).replace(b'\r', b'&#13;')
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)

book = Workbook()
sheet = book.active
sheet.title = 'Data 中文'
sheet.append(['  café 中文 😀  ', '00123', 42.5, True, False])
sheet.append(['=literal', 'tab\tline\nbreak\rreturn & < > "', datetime(2026,10,10,12,30,0,125000), .125, 19.95])
sheet['A2'].data_type = 's'
sheet['C2'].number_format = 'yyyy-mm-dd hh:mm:ss.000'
sheet['D2'].number_format = '0.0%'
sheet['E2'].number_format = '$0.00'
sheet.append([time(6,30,15), timedelta(hours=36,seconds=1), -.125, datetime(1900,2,28), datetime(1900,3,1)])
sheet['B3'].number_format = '[h]:mm:ss.000'
sheet['C3'].number_format = '[h]:mm:ss'
sheet.append(['=SUM(C1,1.5)', '=TRUE()', '="cached text"', '=SUM(1,2)', '#DIV/0!'])
sheet['A5'] = 'Merged anchor'
sheet.merge_cells('A5:B5')
sheet['C6'] = 'https://example.invalid/literal'
sheet['C6'].hyperlink = 'https://example.invalid/never-fetch'
book.create_sheet('Empty')
book.create_sheet('Hidden')['A1'] = 'hidden value'
book['Hidden'].sheet_state = 'hidden'
book.create_sheet('Very hidden')['A1'] = 0
book['Very hidden'].sheet_state = 'veryHidden'
book.save(root/'typed.xlsx')
caches(root/'typed.xlsx', {'A4': ('n','44'), 'B4':('b','1'), 'C4':('str','cached text')})
book = Workbook()
book.epoch = CALENDAR_MAC_1904
book.active.title = 'Epoch 1904'
book.active.append([datetime(1904,1,1),datetime(2026,10,10,12,30),time(0,0),timedelta(hours=49)])
book.active['D1'].number_format='[h]:mm:ss'
book.save(root/'epoch-1904.xlsx')
book = Workbook()
book.active.title = 'Sparse'
book.active['A1'] = 'First cell'
book.active['XFD1048576'] = 'Last cell'
book.save(root/'sparse.xlsx')
print('Generated typed.xlsx, epoch-1904.xlsx, sparse.xlsx')
